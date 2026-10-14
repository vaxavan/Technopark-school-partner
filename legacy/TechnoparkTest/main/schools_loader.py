from pathlib import Path
from urllib.parse import urlparse, urlunparse

from openpyxl import load_workbook


# Папка, в которой находится manage.py
BASE_DIR = Path(__file__).resolve().parent.parent

EXCEL_FILE = BASE_DIR / "schools.xlsx"

SECTION_PATH = "/activity/proforientation/miria"


def load_schools():
    """Загружает список школ-партнёров из Excel."""

    workbook = load_workbook(
        EXCEL_FILE,
        data_only=True
    )

    sheet = workbook.active

    schools = []

    for row in sheet.iter_rows(min_row=2):

        # Столбец B — номер
        number = row[1].value

        # Столбец C — название школы
        name = row[2].value

        # Столбец D — адрес сайта
        website = row[3].value

        if not name or not website:
            continue

        website = str(website).strip()

        # Пропускаем комментарии вместо адресов
        if not website.startswith(("http://", "https://")):
            continue

        parsed = urlparse(website)

        # Пока работаем только с edumsko.ru
        if not (
            parsed.hostname
            and (
                parsed.hostname == "edumsko.ru"
                or parsed.hostname.endswith(".edumsko.ru")
            )
        ):
            continue

        # Берём только домен школы
        base_url = urlunparse((
            parsed.scheme,
            parsed.netloc,
            "",
            "",
            "",
            ""
        ))

        # Формируем единый адрес нужного раздела
        section_url = base_url + SECTION_PATH

        schools.append({
            "number": number,
            "name": str(name).strip(),
            "url": base_url,
            "section_url": section_url
        })

    workbook.close()

    return schools


if __name__ == "__main__":

    print("\nШКОЛЫ-ПАРТНЁРЫ")
    print("=" * 65)

    schools = load_schools()

    for school in schools:

        print(f"\n№ {school['number']}")
        print("Школа:", school["name"])
        print("Сайт:", school["url"])
        print("Раздел:", school["section_url"])

    print("\n" + "=" * 65)
    print("Всего загружено школ:", len(schools))