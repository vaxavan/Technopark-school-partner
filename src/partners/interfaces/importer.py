"""Импорт справочника школ из Excel (формат файла schools.xlsx технопарка).

Ищет строку заголовков, понимает столбцы: территория, №, наименование, сайт,
примечание и (если есть) e-mail. Повторный импорт обновляет школы по названию."""

from urllib.parse import urlparse, urlunparse

from openpyxl import load_workbook

from .. import models as m

EDUMSKO_SECTION = "/activity/proforientation/miria"


def _header_map(row):
    mapping = {}
    for i, cell in enumerate(row):
        v = str(cell or "").strip().lower()
        if not v:
            continue
        if "террит" in v:
            mapping["territory"] = i
        elif v.startswith("№") or v in ("номер", "n"):
            mapping["number"] = i
        elif "наимен" in v or "школ" in v or "название" in v:
            mapping["name"] = i
        elif "сайт" in v or "url" in v:
            mapping["website"] = i
        elif "mail" in v or "почт" in v:
            mapping["email"] = i
        elif "примеч" in v or "коммент" in v:
            mapping["notes"] = i
    return mapping


def default_sources(website):
    """Источники по умолчанию для сайта школы."""
    parsed = urlparse(website)
    base = urlunparse((parsed.scheme, parsed.netloc, "", "", "", ""))
    host = parsed.hostname or ""
    if host.endswith("edumsko.ru"):
        # раздел «Профориентация → РТУ МИРЭА» на школьных сайтах edumsko — проверено в прототипе
        return [("html", base + EDUMSKO_SECTION, EDUMSKO_SECTION + "/post/")]
    if host.endswith("vk.com"):
        return [("vk", website, "")]
    return [("sitemap", base + "/sitemap.xml", ""), ("rss", base + "/rss", "")]


def import_schools(organization, file_obj):
    wb = load_workbook(file_obj, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header_idx, cols = None, {}
    for i, row in enumerate(rows[:15]):
        cols = _header_map(row)
        if "name" in cols:
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("Не найдена строка заголовков (нужен столбец «Наименование»)")
    # в файле технопарка столбец сайта без заголовка — он сразу после названия
    cols.setdefault("website", cols["name"] + 1)
    cols.setdefault("notes", cols["website"] + 1)

    created = updated = 0
    territory = ""
    for row in rows[header_idx + 1:]:
        def get(key):
            i = cols.get(key)
            return row[i] if i is not None and i < len(row) else None

        name = str(get("name") or "").strip()
        if not name:
            continue
        name = " ".join(name.split())
        territory = str(get("territory") or "").strip() or territory
        website = str(get("website") or "").strip()
        notes = str(get("notes") or "").strip()
        email = str(get("email") or "").strip()
        number = get("number")
        is_url = website.startswith(("http://", "https://"))
        if not is_url and website:
            notes = (website + ("; " + notes if notes else "")).strip()

        school, was_created = m.School.objects.get_or_create(organization=organization, name=name)
        school.number = int(number) if isinstance(number, (int, float)) else school.number
        school.territory = territory
        school.notes = notes or school.notes
        if email:
            school.email = email
        if is_url:
            school.website = website
        elif "объедин" in notes.lower():
            school.active = False
        school.save()
        if is_url and not school.sources.exists():
            for kind, url, prefix in default_sources(website):
                m.Source.objects.create(school=school, kind=kind, url=url, link_prefix=prefix)
        created += was_created
        updated += not was_created
    return created, updated
