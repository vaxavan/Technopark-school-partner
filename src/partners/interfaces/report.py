"""Отчёт по школам: по каждой паре «материал — школа», с выгрузкой в Excel."""

from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from publishing.models import Evidence

from ..models import PlacementRecord

STATUS_FILL = {
    "confirmed": "C6EFCE",
    "disputed": "FFEB9C",
    "not_found": "FFC7CE",
    "expected": "FFFFFF",
    "rejected": "D9D9D9",
}


def placement_rows(organization, material=None):
    qs = (PlacementRecord.objects.filter(organization=organization)
          .select_related("school", "material").order_by("material__created_at", "school__number"))
    if material is not None:
        qs = qs.filter(material=material)
    evidence = {(e.material_id, e.school_id): e for e in
                Evidence.objects.filter(organization=organization, school__isnull=False)}
    rows = []
    for p in qs:
        rows.append({"placement": p, "evidence": evidence.get((p.material_id, p.school_id))})
    return rows


def summary(organization, material=None):
    rows = placement_rows(organization, material)
    total = len(rows)
    counts = {}
    for r in rows:
        counts[r["placement"].status] = counts.get(r["placement"].status, 0) + 1
    confirmed = counts.get("confirmed", 0)
    return {"total": total, "counts": counts, "share": round(100 * confirmed / total) if total else 0}


def export_xlsx(organization, material=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Размещение у школ"
    head = ["Материал", "Дата рассылки", "№", "Территория", "Школа", "Результат", "Совпадение, %",
            "Ссылка на публикацию", "Напоминаний", "Проверено", "Доказательство (снимок)", "Решение"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    tz = timezone.get_current_timezone()
    for r in placement_rows(organization, material):
        p, ev = r["placement"], r["evidence"]
        ws.append([
            p.material.title,
            p.dispatched_at.astimezone(tz).strftime("%d.%m.%Y") if p.dispatched_at else "",
            p.school.number,
            p.school.territory,
            p.school.name,
            p.get_status_display(),
            round(p.similarity * 100, 1) if p.similarity is not None else "",
            p.found_url,
            p.reminders_sent,
            p.checked_at.astimezone(tz).strftime("%d.%m.%Y %H:%M") if p.checked_at else "",
            ev.screenshot.url if ev and ev.screenshot else (ev.url if ev else ""),
            p.resolution_note,
        ])
        fill = STATUS_FILL.get(p.status)
        if fill:
            ws.cell(row=ws.max_row, column=6).fill = PatternFill("solid", fgColor=fill)
    widths = [40, 14, 5, 14, 50, 18, 12, 50, 12, 16, 40, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
