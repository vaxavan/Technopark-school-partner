"""Поиск по сайтам школ по теме и выгрузки: Excel «есть / нет» и архив со снимками страниц.

Как это работает:
  1. Администратор вводит фразу («битва роботов»), период и школы — создаётся поиск.
  2. Каждая школа проверяется отдельной фоновой задачей (или агентом обхода, если сайты
     не открываются с сервера): SiteExplorer находит страницы новостей, TopicQuery ищет фразу.
  3. По каждой найденной публикации сохраняются ссылка, заголовок, дата, фрагмент текста,
     снимок страницы и копия HTML с контрольной суммой — как доказательство.
  4. Итог выгружается в Excel и в ZIP: папка на каждую школу со снимками."""

import hashlib
import json
import logging
import re
import tempfile
import zipfile
from datetime import datetime, time
from datetime import timezone as dt_timezone
from io import BytesIO
from urllib.parse import urlparse

from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.events import record
from partners.domain.entities import SourceKind
from partners.domain.topic import TopicQuery
from partners.infrastructure.mappers import source_to_entity
from partners.models import CrawlAgent, School, SchoolSearch, SchoolSearchHit, SchoolSearchResult

log = logging.getLogger(__name__)

RESULT_FILL = {"found": "C6EFCE", "none": "FFC7CE", "unreachable": "FFEB9C", "no_site": "D9D9D9"}


# ---------- запуск

def school_site(school):
    if school.website:
        return school.website
    src = school.sources.exclude(kind=SourceKind.VK.value).first()
    if src:
        u = urlparse(src.url)
        return f"{u.scheme}://{u.netloc}/"
    return ""


def school_sources(school):
    return [source_to_entity(s) for s in school.sources.exclude(kind=SourceKind.VK.value)]


def start_search(org, user, query, since=None, school_ids=None, screenshots=True, on_server=False):
    TopicQuery(query)  # проверка: в запросе есть хотя бы одно слово
    schools = School.objects.filter(organization=org, active=True).prefetch_related("sources")
    if school_ids:
        schools = schools.filter(pk__in=school_ids)
    if not schools.exists():
        raise ValueError("Нет активных школ для поиска")
    agent = CrawlAgent.objects.filter(organization=org).first()
    use_agent = bool(agent and agent.is_active and not on_server)
    with transaction.atomic():
        search = SchoolSearch.objects.create(organization=org, query=query.strip(), since=since,
                                             screenshots=screenshots, created_by=user,
                                             status="agent" if use_agent else "queued")
        for school in schools:
            has_site = bool(school_site(school))
            SchoolSearchResult.objects.create(search=search, school=school,
                                              status="queued" if has_site else "no_site",
                                              finished_at=None if has_site else timezone.now())
        record(org, "search.started", actor=user, obj=search,
               message=f"Поиск по сайтам школ: «{search.query}», школ: {schools.count()}"
                       + (" (через агента обхода)" if use_agent else ""))
    if not use_agent:
        from partners.tasks import run_search

        transaction.on_commit(lambda: run_search.delay(search.pk))
    return search


def queue_on_server(search, only_unreachable=False):
    """Проверить на сервере школы, которые ещё ждут (или сайт не открылся у агента)."""
    statuses = ["unreachable"] if only_unreachable else ["queued", "unreachable"]
    search.results.filter(status__in=statuses).update(status="queued", error="", finished_at=None)
    search.status = "queued"
    search.finished_at = None
    search.save(update_fields=["status", "finished_at"])
    from partners.tasks import run_search

    transaction.on_commit(lambda: run_search.delay(search.pk))


def cancel(search):
    search.results.filter(status="queued").update(status="none", error="поиск остановлен",
                                                  finished_at=timezone.now())
    search.status = "cancelled"
    search.finished_at = timezone.now()
    search.save(update_fields=["status", "finished_at"])


def since_dt(search):
    return datetime.combine(search.since, time.min, tzinfo=dt_timezone.utc) if search.since else None


# ---------- выполнение (сервер)

def run_school(result_id, explorer=None, snapshotter=None):
    from integrations import snapshot as snapshots
    from partners.infrastructure.explorer import SiteExplorer

    claimed = SchoolSearchResult.objects.filter(pk=result_id, status="queued").update(
        status="running", started_at=timezone.now())
    if not claimed:
        return None
    result = SchoolSearchResult.objects.select_related("search__organization", "school").get(pk=result_id)
    search, school = result.search, result.school
    if search.status == "cancelled":
        return _finish(result, "none", error="поиск остановлен")
    query = TopicQuery(search.query)
    explorer = explorer or SiteExplorer()
    snapshotter = (snapshotter or snapshots.get_snapshotter()) if search.screenshots else None
    try:
        report = explorer.explore(school_site(school), school_sources(school), query, since_dt(search))
    except Exception as exc:  # одна школа не должна останавливать поиск
        log.exception("Поиск по сайту школы %s прервался", school.pk)
        return _finish(result, "unreachable", error=str(exc)[:1000])
    hits = 0
    for page in report.pages:
        m = query.match(page.title, page.text)
        if m is None:
            continue
        shot = None
        if snapshotter is not None:
            snap = snapshotter.capture(page.url)
            shot = snap.png if snap else None
        save_hit(result, url=page.url, title=page.title, published_at=page.published_at, snippet=m.snippet,
                 score=m.score, html=page.html, screenshot=shot)
        hits += 1
    result.pages_read = report.read
    if hits:
        status = "found"
    elif report.unreachable or (not report.pages and report.errors):
        status = "unreachable"
    else:
        status = "none"
    return _finish(result, status, error="; ".join(report.errors[:5])[:2000] if status != "found" else "")


def save_hit(result, *, url, title, published_at, snippet, score, html="", screenshot=None):
    hit, _ = SchoolSearchHit.objects.update_or_create(
        result=result, url=url[:700],
        defaults={"title": (title or "")[:500], "published_at": published_at, "snippet": snippet or "",
                  "score": score or 0, "captured_at": timezone.now()})
    n = result.hits.count()
    stamp = (published_at or timezone.now()).strftime("%Y-%m-%d")
    if html:
        data = html.encode("utf-8") if isinstance(html, str) else html
        hit.checksum = hashlib.sha256(data).hexdigest()
        hit.html_copy.save(f"{stamp}-{n}.html", ContentFile(data), save=False)
    if screenshot:
        ext = "jpg" if screenshot[:3] == b"\xff\xd8\xff" else "png"
        hit.screenshot.save(f"{stamp}-{n}.{ext}", ContentFile(screenshot), save=False)
    hit.save()
    return hit


def _finish(result, status, error=""):
    result.status, result.error = status, error
    result.hits_count = result.hits.count()
    result.finished_at = timezone.now()
    result.save(update_fields=["status", "error", "hits_count", "pages_read", "finished_at"])
    maybe_finish(result.search)
    return result


def maybe_finish(search):
    search.refresh_from_db(fields=["status"])
    if search.status in ("done", "cancelled"):
        return
    if search.results.filter(status__in=["queued", "running"]).exists():
        if search.status == "queued":
            SchoolSearch.objects.filter(pk=search.pk, status="queued").update(status="running")
        return
    updated = SchoolSearch.objects.filter(pk=search.pk).exclude(status__in=["done", "cancelled"]).update(
        status="done", finished_at=timezone.now())
    if updated:
        counts = summary(search)
        record(search.organization, "search.done", obj=search,
               message=f"Поиск «{search.query}» завершён: есть у {counts['found']}, нет у {counts['none']}, "
                       f"сайт не открылся у {counts['unreachable']}")


def summary(search):
    counts = {k: 0 for k, _ in SchoolSearchResult.STATUS_CHOICES}
    for status in search.results.values_list("status", flat=True):
        counts[status] = counts.get(status, 0) + 1
    total = sum(counts.values())
    done = total - counts["queued"] - counts["running"]
    return {**counts, "total": total, "done": done, "percent": round(100 * done / total) if total else 100}


# ---------- результаты от агента обхода

def save_agent_result(result, data, files):
    """Агент проверил сайт школы у себя: data — итог и найденные публикации, files — снимки и копии."""
    hits = data.get("hits")
    if isinstance(hits, str):
        hits = json.loads(hits or "[]")
    for i, h in enumerate(hits or []):
        url = str(h.get("url", ""))
        if not url.startswith(("http://", "https://")):
            continue
        published = h.get("published_at")
        try:
            published = datetime.fromisoformat(published) if published else None
        except (TypeError, ValueError):
            published = None
        shot = files.get(f"shot_{i}")
        html = files.get(f"html_{i}")
        save_hit(result, url=url, title=str(h.get("title", "")), published_at=published,
                 snippet=str(h.get("snippet", ""))[:2000], score=float(h.get("score") or 0),
                 html=html.read() if html else "", screenshot=shot.read() if shot else None)
    status = data.get("status")
    if status not in ("found", "none", "unreachable"):
        status = "found" if result.hits.exists() else "none"
    result.pages_read = int(data.get("pages_read") or 0)
    result.checked_by_agent = True
    result.save(update_fields=["pages_read", "checked_by_agent"])
    return _finish(result, status, error=str(data.get("error", ""))[:2000])


# ---------- выгрузки

def _safe(name, limit=60):
    name = re.sub(r'[\\/:*?"<>|\n\r\t]+', " ", name or "").strip()
    return re.sub(r"\s+", " ", name)[:limit].strip() or "без названия"


def _school_folder(school):
    num = f"{school.number:02d} " if school.number else ""
    return _safe(num + school.name, 70)


def hit_files(hit):
    """Имена файлов доказательства внутри архива: папка школы / дата заголовок.png|html."""
    date = timezone.localtime(hit.published_at).strftime("%Y-%m-%d") if hit.published_at else "без-даты"
    base = f"{_school_folder(hit.result.school)}/{date} {_safe(hit.title, 50)} ({hit.pk})"
    shot = f"{base}.{hit.screenshot.name.rsplit('.', 1)[-1]}" if hit.screenshot else ""
    html = f"{base}.html" if hit.html_copy else ""
    return shot, html


def _results(search):
    return (search.results.select_related("school").prefetch_related("hits")
            .order_by("school__number", "school__name"))


def export_search_xlsx(search) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Школы"
    ws.append([f"Поиск «{search.query}»" + (f", публикации с {search.since:%d.%m.%Y}" if search.since else "")])
    ws["A1"].font = Font(bold=True, size=12)
    head = ["№", "Территория", "Школа", "Сайт", "Есть публикации?", "Сколько", "Самая ранняя публикация", "Её дата",
            "Проверено", "Примечание"]
    ws.append(head)
    for c in ws[2]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    labels = {"found": "Есть", "none": "Нет", "unreachable": "Сайт не открылся", "no_site": "Нет адреса сайта",
              "queued": "Ещё не проверено", "running": "Проверяется"}
    for r in _results(search):
        hits = list(r.hits.all())
        first = min(hits, key=lambda h: h.published_at or timezone.now()) if hits else None
        ws.append([r.school.number, r.school.territory, r.school.name, school_site(r.school), labels[r.status],
                   len(hits), first.url if first else "",
                   timezone.localtime(first.published_at).strftime("%d.%m.%Y") if first and first.published_at else "",
                   timezone.localtime(r.finished_at).strftime("%d.%m.%Y %H:%M") if r.finished_at else "",
                   ("проверил агент обхода; " if r.checked_by_agent else "") + (r.error[:200] if r.error else "")])
        fill = RESULT_FILL.get(r.status)
        if fill:
            ws.cell(row=ws.max_row, column=5).fill = PatternFill("solid", fgColor=fill)
    for i, w in enumerate([5, 14, 50, 32, 18, 9, 50, 11, 16, 40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"

    ws2 = wb.create_sheet("Публикации")
    ws2.append(["Школа", "Заголовок", "Дата", "Ссылка", "Фрагмент", "Совпадение", "Снимок в архиве",
                "Копия страницы в архиве", "SHA-256 копии"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for r in _results(search):
        for h in r.hits.all():
            shot, html = hit_files(h)
            ws2.append([r.school.name, h.title,
                        timezone.localtime(h.published_at).strftime("%d.%m.%Y") if h.published_at else "",
                        h.url, h.snippet, f"{round(h.score * 100)}%", shot, html, h.checksum])
            ws2.cell(row=ws2.max_row, column=5).alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate([40, 50, 11, 50, 70, 11, 45, 45, 20], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _readme(search):
    return (f"Поиск по сайтам школ: «{search.query}»\n"
            f"Создан: {timezone.localtime(search.created_at):%d.%m.%Y %H:%M}\n"
            + (f"Публикации не раньше: {search.since:%d.%m.%Y}\n" if search.since else "")
            + "\nИтоги.xlsx — по каждой школе: есть ли публикации, ссылки, даты.\n"
              "Папки — по школам: снимок страницы (.png/.jpg) и копия страницы (.html) на момент проверки.\n"
              "Контрольная сумма SHA-256 копии указана в Итоги.xlsx, лист «Публикации».\n")


def export_search_zip(search):
    """ZIP во временном файле (не держим архив в памяти). Возвращает открытый файл."""
    tmp = tempfile.TemporaryFile()
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("Итоги.xlsx", export_search_xlsx(search))
        z.writestr("ПРОЧТИТЕ.txt", _readme(search))
        for r in _results(search):
            for h in r.hits.all():
                shot, html = hit_files(h)
                for name, f in ((shot, h.screenshot), (html, h.html_copy)):
                    if name and f:
                        try:
                            with f.open("rb") as fh:
                                z.writestr(name, fh.read())
                        except OSError:
                            log.warning("Файл доказательства не найден: %s", f.name)
    tmp.seek(0)
    return tmp


def export_evidence_zip(organization, material=None):
    """Архив доказательств размещения наших материалов: папка на материал, внутри — по школам."""
    from partners.interfaces.report import export_xlsx
    from publishing.models import Evidence

    tmp = tempfile.TemporaryFile()
    qs = Evidence.objects.filter(organization=organization, school__isnull=False).select_related("material", "school")
    if material is not None:
        qs = qs.filter(material=material)
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("Отчёт по школам.xlsx", export_xlsx(organization, material))
        for e in qs:
            folder = f"{_safe(e.material.title, 50)}/{_school_folder(e.school)}"
            stamp = timezone.localtime(e.captured_at).strftime("%Y-%m-%d %H-%M")
            for f in (e.screenshot, e.html_copy):
                if f:
                    try:
                        with f.open("rb") as fh:
                            z.writestr(f"{folder}/{stamp}.{f.name.rsplit('.', 1)[-1]}", fh.read())
                    except OSError:
                        log.warning("Файл доказательства не найден: %s", f.name)
            z.writestr(f"{folder}/ссылка.txt", f"{e.url}\nСнято: {stamp}\nSHA-256 копии: {e.checksum}\n")
    tmp.seek(0)
    return tmp
