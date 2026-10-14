"""Весь процесс только через REST API, с проверкой CSRF — так будет работать отдельный фронтенд.
Если этот тест проходит, React-клиенту не нужно ничего, кроме API."""

import io
from pathlib import Path

import pytest
from django.test import Client
from openpyxl import load_workbook

from materials.models import Channel


def api(client, method, url, data=None, **kw):
    token = client.cookies["csrftoken"].value if "csrftoken" in client.cookies else ""
    fn = getattr(client, method)
    if method == "get":
        return fn(url, data or {})
    if kw.get("multipart"):
        return fn(url, data or {}, HTTP_X_CSRFTOKEN=token)
    return fn(url, data or {}, content_type="application/json", HTTP_X_CSRFTOKEN=token)


@pytest.mark.django_db
def test_whole_process_through_api(org, admin_user, director):
    from partners.models import School

    School.objects.create(organization=org, name="СОШ №16", email="s16@example.com")

    # --- автор, без входа ---
    author = Client(enforce_csrf_checks=True)
    assert api(author, "get", "/api/v1/public/istok/").json()["name"] == org.name
    r = api(author, "post", "/api/v1/public/istok/submissions/", {
        "title": "Конкурс роботов", "text": "Прошёл конкурс роботов, участвовали 8 команд.",
        "place": "Фрязино", "author_name": "Анна", "author_email": "anna@example.com", "consent": "true",
    }, multipart=True)
    assert r.status_code == 201, r.content
    sub = r.json()
    assert sub["finished"] is False and sub["round"]["questions"][0]["type"] == "text"
    r = api(author, "post", f"/api/v1/public/istok/submissions/{sub['id']}/questions/",
            {"answers": {"q1": "12 октября"}})
    assert r.json()["finished"] is True

    # --- администратор ---
    staff = Client(enforce_csrf_checks=True)
    api(staff, "get", "/api/v1/auth/csrf/")
    assert api(staff, "post", "/api/v1/auth/login/", {"email": "admin@example.com", "password": "wrong"}).status_code == 400
    me = api(staff, "post", "/api/v1/auth/login/", {"email": "admin@example.com", "password": "pass12345"}).json()
    assert me["organizations"][0]["role"] == "admin"
    api(staff, "get", "/api/v1/auth/csrf/")  # после входа токен обновляется

    # без CSRF-заголовка изменение отклоняется
    assert staff.post(f"/api/v1/submissions/{sub['id']}/accept/", {}, content_type="application/json").status_code == 403

    dash = api(staff, "get", "/api/v1/dashboard/").json()
    assert dash["counts"]["new"] == 1 and dash["new_submissions"][0]["id"] == sub["id"]

    r = api(staff, "post", f"/api/v1/submissions/{sub['id']}/accept/", {"channels": ["vk", "dzen", "partners"]})
    mid = r.json()["id"]
    m = api(staff, "get", f"/api/v1/materials/{mid}/").json()
    assert m["status"] == "on_approval" and {v["channel"] for v in m["variants"]} == {"vk", "dzen", "partners"}

    v = [x for x in m["variants"] if x["channel"] == "vk"][0]
    api(staff, "post", f"/api/v1/materials/{mid}/edit/", {"channel": "vk", "text": v["text"] + " Спасибо всем!"})
    hist = api(staff, "get", f"/api/v1/materials/{mid}/history/").json()
    first = [h for h in hist if h["channel"] == "vk" and h["version"] == 1][0]
    api(staff, "post", f"/api/v1/materials/{mid}/restore/", {"variant_id": first["id"]})
    api(staff, "post", f"/api/v1/materials/{mid}/channels/", {"channels": ["vk", "dzen", "partners", "max"]})
    m = api(staff, "get", f"/api/v1/materials/{mid}/").json()
    assert "max" in {x["channel"] for x in m["variants"]}

    r = api(staff, "post", f"/api/v1/materials/{mid}/approve/", {"basis": "oral", "comment": "директор в отъезде"})
    assert r.json()["status"] == "scheduled"

    from publishing.tasks import publish_due

    assert publish_due() == 3
    pubs = api(staff, "get", "/api/v1/publications/").json()
    assert {p["channel"] for p in pubs} == {"vk", "dzen", "max"} and all(p["status"] == "published" for p in pubs)

    rep = api(staff, "get", f"/api/v1/materials/{mid}/placements/").json()
    assert rep["summary"]["total"] == 1
    xlsx = api(staff, "get", "/api/v1/placements/export/", {"material": mid})
    assert load_workbook(io.BytesIO(xlsx.content)).active.max_row == 2

    # школы: импорт и проверка
    path = Path(__file__).resolve().parent.parent / "data" / "schools.xlsx"
    with open(path, "rb") as fh:
        r = api(staff, "post", "/api/v1/schools/import/", {"file": fh}, multipart=True)
    assert r.json()["created"] == 18
    assert api(staff, "get", "/api/v1/schools/", {"q": "Лицей"}).json()["count"] >= 2

    # настройки: токен уходит в базу зашифрованным и наружу не отдаётся
    s = api(staff, "put", "/api/v1/settings/platforms/vk/", {"target_id": "123", "token": "vk1.a.SECRET"}).json()
    assert s["platforms"]["vk"] == {"target_id": "123", "enabled": True, "has_token": True}
    assert "SECRET" not in str(api(staff, "get", "/api/v1/settings/").content)
    api(staff, "patch", "/api/v1/settings/", {"style_guide": "Пишем коротко."})
    assert api(staff, "get", "/api/v1/settings/").json()["organization"]["style_guide"] == "Пишем коротко."

    # школа присылает ссылку по токену из письма
    from partners.models import PlacementRecord

    token = PlacementRecord.objects.get().report_token
    school = Client(enforce_csrf_checks=True)
    assert api(school, "get", f"/api/v1/public/school-link/{token}/").json()["school"] == "СОШ №16"

    assert api(staff, "post", "/api/v1/auth/logout/").status_code == 204
    assert staff.get("/api/v1/dashboard/").status_code == 403


@pytest.mark.django_db
def test_director_sees_basis_options(org, director, submission):
    c = Client()
    c.force_login(director)
    from materials import services

    m = services.accept_submission(submission, director, [Channel.VK])
    r = c.post(f"/api/v1/materials/{m.public_id}/approve/", {"basis": "director"}, content_type="application/json")
    assert r.status_code == 200 and r.json()["approvals"][0]["basis"] == "director"


def test_openapi_schema_builds(db):
    r = Client().get("/api/schema/")
    assert r.status_code == 200 and b"/api/v1/public/{code}/submissions/" in r.content
