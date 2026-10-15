"""Сквозной сценарий: заявка → уточнения → принятие → тексты → утверждение → публикация →
доказательство → рассылка школам → обнаружение размещения → отчёт."""

import io
import json

import pytest
from django.core import mail
from django.urls import reverse
from openpyxl import load_workbook

from core.events import deliver_pending
from materials.models import Channel, Material
from partners.domain.entities import FoundPublication, SourceKind
from partners.models import PlacementRecord, School, Source
from publishing.models import Evidence, Publication
from submissions.models import Submission


@pytest.mark.django_db
def test_full_flow(client, org, admin_user, director, monkeypatch):
    # школы-партнёры
    s1 = School.objects.create(organization=org, name="СОШ №16", email="s16@example.com")
    Source.objects.create(school=s1, kind="html", url="https://s16.example/news")
    s2 = School.objects.create(organization=org, name="Лицей №14", email="l14@example.com")
    Source.objects.create(school=s2, kind="html", url="https://l14.example/news")
    School.objects.create(organization=org, name="Без почты")

    # 1. автор отправляет заявку без даты — модель задаёт вопрос
    resp = client.post(reverse("public_form", args=[org.code]), {
        "title": "Конкурс роботов", "text": "Прошёл конкурс роботов в технопарке, участвовали 8 команд.",
        "place": "Фрязино", "author_name": "Анна", "author_email": "anna@example.com", "consent": "on"})
    assert resp.status_code == 302
    sub = Submission.objects.get()
    questions_url = resp["Location"]
    # сначала экран ожидания: модель проверяет заявку, страница сама спрашивает «что дальше»
    page = client.get(questions_url).content.decode()
    assert "think-art" in page and sub.rounds.count() == 0
    think_url = reverse("public_think", args=[org.code, sub.public_id])
    r = client.post(think_url, HTTP_ACCEPT="application/json")
    assert r.json() == {"state": "questions", "next": questions_url}
    resp = client.get(questions_url)
    assert "Когда проходило мероприятие?" in resp.content.decode()
    assert sub.rounds.count() == 1

    # 2. автор отвечает; при следующем раунде сведений достаточно
    resp = client.post(questions_url, {"q1": "12 октября", "q2": "да", "answer": "1"})
    assert "think-art" in client.get(resp["Location"]).content.decode()
    resp = client.post(think_url, follow=True)          # без JavaScript — обычная кнопка «Продолжить»
    assert "принята" in resp.content.decode()
    sub.refresh_from_db()
    assert sub.questions_finished and "12 октября" in sub.full_text_for_model()
    deliver_pending()
    to = [m.to[0] for m in mail.outbox]
    assert "anna@example.com" in to and "admin@example.com" in to

    # 3. администратор принимает заявку — тексты готовятся (фоновая задача, в тестах сразу)
    client.force_login(admin_user)
    resp = client.post(reverse("submission_detail", args=[sub.public_id]),
                       {"action": "accept", "channels": [Channel.VK, Channel.DZEN, Channel.PARTNERS]})
    material = Material.objects.get()
    assert resp["Location"] == reverse("material_detail", args=[material.public_id])
    material.refresh_from_db()
    assert material.status == "on_approval"
    assert set(material.current_variants().values_list("channel", flat=True)) == {"vk", "dzen", "partners"}

    # 4. ручная правка создаёт новую версию и считает долю ручных правок
    v1 = material.current_variant(Channel.VK)
    client.post(reverse("material_detail", args=[material.public_id]),
                {"action": "save_text", "channel": "vk", "text": v1.text + "\n\n#технопарк #роботы"})
    v2 = material.current_variant(Channel.VK)
    assert v2.version == 2 and v2.source == "manual" and v2.manual_share > 0

    # 5. администратор не может поставить «утвердил директор»
    resp = client.post(reverse("material_detail", args=[material.public_id]),
                       {"action": "approve", "basis": "director"}, follow=True)
    material.refresh_from_db()
    assert material.status == "on_approval"

    # утверждает с основанием «устно согласовано» — директор получает уведомление
    mail.outbox.clear()
    client.post(reverse("material_detail", args=[material.public_id]), {"action": "approve", "basis": "oral"})
    material.refresh_from_db()
    assert material.status == "scheduled"
    assert Publication.objects.filter(material=material).count() == 2  # vk + dzen; партнёры — рассылка
    deliver_pending()
    assert any(m.to == ["director@example.com"] for m in mail.outbox)

    # 6. планировщик публикует (без токенов — демо), повторный запуск ничего не дублирует
    from publishing.tasks import publish_due

    mail.outbox.clear()
    assert publish_due() == 2
    assert publish_due() == 0
    material.refresh_from_db()
    assert material.status == "published"
    vk = material.publications.get(channel="vk")
    assert vk.demo and vk.url.endswith(f"/demo/vk/{material.public_id}/")
    assert Evidence.objects.filter(publication=vk).exclude(checksum="").exists()

    # публикация в Дзен = элемент RSS-ленты
    rss = client.get(reverse("rss_feed", args=[org.code])).content.decode()
    assert "<item>" in rss and str(material.public_id) in rss

    # 7. рассылка школам ушла после публикации, школе без почты — нет
    assert PlacementRecord.objects.filter(material=material).count() == 3
    school_mails = [m for m in mail.outbox if m.to[0] in ("s16@example.com", "l14@example.com")]
    assert len(school_mails) == 2 and "List-Unsubscribe" in school_mails[0].extra_headers

    # 8. обход сайтов: у школы 16 материал найден, у лицея нет
    text = material.current_variant(Channel.PARTNERS).text

    class FakeCrawler:
        def fetch(self, source):
            if "s16" in source.url:
                return [FoundPublication("https://s16.example/news/1", "Конкурс", "Новости школы. " + text,
                                         SourceKind.HTML)]
            return []

    from partners import container

    report = container.detect_placement(FakeCrawler())(org.id)
    assert report.confirmed == 1
    p16 = PlacementRecord.objects.get(school=s1)
    assert p16.status == "confirmed" and p16.found_url == "https://s16.example/news/1"
    assert Evidence.objects.filter(school=s1, material=material).exists()

    # 9. отчёт и выгрузка в Excel
    resp = client.get(reverse("partners_report") + f"?m={material.public_id}&export=1")
    ws = load_workbook(io.BytesIO(resp.content)).active
    statuses = {row[4]: row[5] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert statuses["СОШ №16"] == "Подтверждено" and statuses["Лицей №14"] == "Ожидается"

    # 10. все экраны открываются
    for name, args in [("dashboard", []), ("submission_list", []), ("material_list", []), ("calendar", []),
                       ("school_list", []), ("partners_report", []), ("events", []), ("settings", []),
                       ("material_detail", [material.public_id]), ("submission_detail", [sub.public_id]),
                       ("school_edit", [s1.pk])]:
        assert client.get(reverse(name, args=args)).status_code == 200, name
    assert client.get(reverse("news_page", args=[org.code, material.public_id])).status_code == 200
    assert client.get(reverse("demo_post", args=["vk", material.public_id])).status_code == 200


@pytest.mark.django_db
def test_school_reports_link_itself(client, org, admin_user, submission, monkeypatch):
    from materials import services
    from partners.infrastructure import crawlers

    material = services.accept_submission(submission, admin_user, [Channel.PARTNERS])
    school = School.objects.create(organization=org, name="СОШ №1", email="s1@example.com")
    services.approve(material, admin_user, "none")
    p = PlacementRecord.objects.get(school=school)
    text = material.current_variant(Channel.PARTNERS).text
    monkeypatch.setattr(crawlers.HttpClient, "get", lambda self, url: f"<html><h1>Новость</h1><article>{text}</article></html>")
    resp = client.post(reverse("school_report", args=[p.report_token]), {"url": "https://s1.example/n/5"})
    assert resp.status_code == 200
    p.refresh_from_db()
    assert p.status == "confirmed" and p.found_url == "https://s1.example/n/5"

    client.post(reverse("school_unsubscribe", args=[p.report_token]))
    school.refresh_from_db()
    assert school.unsubscribed_at is not None


@pytest.mark.django_db
def test_author_adds_info(client, org, admin_user, submission):
    from materials import services

    services.request_info(submission, admin_user, "Пришлите фото")
    resp = client.post(reverse("public_add", args=[org.code, submission.public_id]), {"text": "Было 40 зрителей"})
    assert resp.status_code == 200
    submission.refresh_from_db()
    assert submission.status == "new" and "40 зрителей" in submission.text


@pytest.mark.django_db
def test_rate_limit_public_form(client, org):
    data = {"title": "т", "text": "т", "author_name": "а", "author_email": "a@example.com", "consent": "on"}
    codes = [client.post(reverse("public_form", args=[org.code]), data).status_code for _ in range(6)]
    assert codes[:5] == [302] * 5 and codes[5] == 429


@pytest.mark.django_db
def test_api_flow_and_isolation(client, org, other_org, admin_user, submission, django_user_model):
    client.force_login(admin_user)
    r = client.get("/api/v1/submissions/")
    assert r.status_code == 200 and r.json()["count"] == 1
    r = client.post(f"/api/v1/submissions/{submission.public_id}/accept/", {"channels": ["vk"]},
                    content_type="application/json")
    assert r.status_code == 201
    mid = r.json()["id"]
    r = client.post(f"/api/v1/materials/{mid}/edit/", {"channel": "vk", "command": "fix"},
                    content_type="application/json")
    assert r.status_code == 200
    r = client.post(f"/api/v1/materials/{mid}/approve/", {"basis": "director"}, content_type="application/json")
    assert r.status_code == 400 and "директор" in json.dumps(r.json(), ensure_ascii=False)

    # пользователь другой организации ничего не видит
    from accounts.models import Membership

    stranger = django_user_model.objects.create_user("stranger@example.com", password="pass12345")
    Membership.objects.create(user=stranger, organization=other_org, role="admin")
    client.force_login(stranger)
    assert client.get(f"/api/v1/materials/{mid}/").status_code == 404
    assert client.get(reverse("material_detail", args=[mid])).status_code == 404
    assert client.get("/api/v1/submissions/").json()["count"] == 0

    client.logout()
    assert client.get("/api/v1/submissions/").status_code == 403


@pytest.mark.django_db
def test_login_by_email_any_case(client, admin_user):
    resp = client.post(reverse("login"), {"username": "  ADMIN@Example.com ", "password": "pass12345"})
    assert resp.status_code == 302 and resp["Location"] == reverse("dashboard")


@pytest.mark.django_db
def test_init_system_creates_users_and_schools(monkeypatch, django_user_model):
    from django.core.management import call_command

    from accounts.models import Membership

    monkeypatch.setenv("DIRECTOR_EMAIL", "Boss@Technopark.ru")
    monkeypatch.setenv("DIRECTOR_PASSWORD", "secret-pass-1")
    monkeypatch.delenv("ADMIN_EMAIL", raising=False)
    call_command("init_system")
    call_command("init_system")  # повторный запуск ничего не дублирует
    assert django_user_model.objects.count() == 2
    boss = django_user_model.objects.get(email="boss@technopark.ru")
    assert boss.check_password("secret-pass-1")
    assert Membership.objects.get(user=boss).role == "director"
    assert django_user_model.objects.filter(email="admin@istok.local").exists()
    assert boss.memberships.get().organization.schools.count() == 18


@pytest.mark.django_db
def test_double_submit_and_back_button_make_one_submission(client, org):
    """Двойное нажатие, «назад» и повторная отправка — одна заявка, а не несколько."""
    page = client.get(reverse("public_form", args=[org.code])).content.decode()
    token = page.split('name="form_token" value="')[1].split('"')[0]
    data = {"title": "Вторая новость", "text": "Прошёл турнир.", "author_name": "Иван",
            "author_email": "ivan@example.com", "consent": "on", "form_token": token}
    first = client.post(reverse("public_form", args=[org.code]), data)
    second = client.post(reverse("public_form", args=[org.code]), data)
    assert first["Location"] == second["Location"]
    # та же новость из новой вкладки (другой номер формы) в течение получаса — тоже не дубль
    third = client.post(reverse("public_form", args=[org.code]), {**data, "form_token": ""})
    assert third["Location"] == first["Location"]
    assert Submission.objects.count() == 1
    # а другая новость того же автора — новая заявка
    client.post(reverse("public_form", args=[org.code]), {**data, "form_token": "", "title": "Третья новость"})
    assert Submission.objects.count() == 2


@pytest.mark.django_db
def test_two_rounds_max_and_staff_notified_once(client, org, admin_user, monkeypatch):
    from submissions import services as sub_services

    class AlwaysAsks:
        def complete(self, system, user, **kw):
            return '{"complete": false, "missing": ["итог"], "questions": [{"id": "q1", "type": "text", "text": "Кто победил?"}]}'

    monkeypatch.setattr("submissions.clarify.get_llm", lambda org=None: AlwaysAsks())
    sub, _ = sub_services.create_submission(org, {"title": "Турнир", "text": "Прошёл турнир.", "author_name": "Иван",
                                                  "author_email": "i@example.com"})
    assert sub_services.prepare_round(sub) == sub_services.QUESTIONS
    assert sub_services.answer_round(sub, {"q1": "Школа №1"}) == sub_services.THINKING
    assert sub_services.prepare_round(sub) == sub_services.QUESTIONS     # второй раунд
    assert sub_services.answer_round(sub, {"q1": "Команда 5Б"}) == sub_services.DONE   # третьего нет
    assert sub.rounds.count() == 2
    sub_services.finish_submission(sub)
    from core.models import Event
    assert Event.objects.filter(kind="submission.received").count() == 1


@pytest.mark.django_db
def test_parallel_think_requests_ask_model_once(org, monkeypatch):
    """Вторая вкладка или повторный запрос, пока модель думает, не создают второй раунд."""
    from submissions import services as sub_services
    from submissions.models import Submission as S

    calls = []

    class Slow:
        def complete(self, system, user, **kw):
            calls.append(1)
            # пока «думаем», приходит второй запрос — он должен просто подождать
            assert sub_services.prepare_round(S.objects.get(pk=sub.pk)) == sub_services.THINKING
            return '{"complete": false, "missing": [], "questions": [{"id": "q1", "type": "text", "text": "Где?"}]}'

    monkeypatch.setattr("submissions.clarify.get_llm", lambda org=None: Slow())
    sub, _ = sub_services.create_submission(org, {"title": "Т", "text": "Т", "author_name": "И",
                                                  "author_email": "i@example.com"})
    assert sub_services.prepare_round(sub) == sub_services.QUESTIONS
    assert len(calls) == 1 and sub.rounds.count() == 1


@pytest.mark.django_db
def test_stuck_model_and_abandoned_author_do_not_lose_submission(org):
    from datetime import timedelta

    from django.utils import timezone

    from core.models import Event
    from submissions import services as sub_services
    from submissions.models import Submission as S

    stuck, _ = sub_services.create_submission(org, {"title": "А", "text": "А", "author_name": "И",
                                                    "author_email": "a@example.com"})
    S.objects.filter(pk=stuck.pk).update(clarify_state="running",
                                         clarify_started_at=timezone.now() - timedelta(minutes=5))
    assert sub_services.clarify_state(stuck) == sub_services.DONE          # модель зависла — идём дальше

    gone, _ = sub_services.create_submission(org, {"title": "Б", "text": "Б", "author_name": "И",
                                                   "author_email": "b@example.com"})
    S.objects.filter(pk=gone.pk).update(updated_at=timezone.now() - timedelta(hours=1))
    assert sub_services.finish_abandoned() == 1
    gone.refresh_from_db()
    assert gone.questions_finished
    assert Event.objects.filter(kind="submission.received").count() == 2
