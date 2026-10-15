"""Поиск по сайтам школ по теме: обход сайта, совпадения, снимки, выгрузки, работа через агента."""

import io

import zipfile
from datetime import date, datetime, timezone

import pytest
from openpyxl import load_workbook

from partners.domain.topic import TopicQuery
from partners.infrastructure.crawlers import HttpClient, SourceUnavailable
from partners.infrastructure.explorer import SiteExplorer
from partners.interfaces import search as svc
from partners.models import CrawlAgent, School, SchoolSearch, Source

FILLER = " Ребята работали командами, наставники помогали, все получили грамоты." * 8


def page(title, body, date_text=""):
    return (f'<html><head><title>{title}</title></head><body><h1>{title}</h1>'
            f'<div class="user_content"><p>{date_text} {body}{FILLER}</p></div></body></html>')


SITES = {
    # школа 16: новости на edumsko-подобном сайте, нужная — на второй странице списка
    "https://s16.example/robots.txt": "User-agent: *",
    "https://s16.example/": '<a href="/press/news">Новости школы</a><a href="/about">О нас</a>',
    "https://s16.example/press/news": '<a href="/press/news/post/12">Осенний бал</a>'
                                      '<a href="/press/news?page=2">2</a>',
    "https://s16.example/press/news?page=2": '<a href="/press/news/post/7">Наши победы</a>'
                                             '<a href="/press/news/post/3">Старая битва роботов</a>',
    "https://s16.example/press/news/post/12": page("Осенний бал", "Прошёл осенний бал.", "10 октября 2026"),
    "https://s16.example/press/news/post/7": page("Наши победы", "Команда школы победила в битве роботов "
                                                  "в технопарке «Исток».", "12 октября 2026"),
    "https://s16.example/press/news/post/3": page("Старая битва роботов", "Битва роботов прошла в спортзале.",
                                                  "5 марта 2024"),
    # лицей 14: сайт открывается, про роботов ничего нет
    "https://l14.example/robots.txt": "Sitemap: https://l14.example/sm.xml",
    "https://l14.example/sm.xml": '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">'
                                  '<url><loc>https://l14.example/news/1</loc><lastmod>2026-09-01</lastmod></url></urlset>',
    "https://l14.example/": "<p>Главная</p>",
    "https://l14.example/news/1": page("Экскурсия", "Экскурсия в музей.", "1 сентября 2026"),
}


@pytest.fixture
def fake_web(monkeypatch):
    calls = []

    def get(self, url):
        calls.append(url)
        if "down.example" in url:
            raise SourceUnavailable(f"{url}: сайт не отвечает на подключение (время ожидания истекло)")
        if url in SITES:
            return SITES[url]
        raise SourceUnavailable(f"{url}: сайт ответил кодом 404")

    monkeypatch.setattr(HttpClient, "get", get)
    return calls


def test_topic_query_forms_and_false_friends():
    q = TopicQuery("битва роботов")
    assert q.match("В битве роботов победила команда", "")
    assert q.match("Роботов собрали, битва была жаркой", "")
    assert not q.match("Робототехника и битвы за урожай", "")
    assert not q.match("Новости", "Битва за урожай закончилась, все довольны и отдыхают. Через неделю "
                                  "в школу приехали роботы.")                        # слова далеко друг от друга
    many = TopicQuery("битва роботов, технопарк Исток")
    assert many.match("Экскурсия в технопарк «Исток»", "").phrase == "технопарк исток"


def test_explorer_follows_pagination_and_period(fake_web):
    since = datetime(2026, 1, 1, tzinfo=timezone.utc)
    report = SiteExplorer(HttpClient()).explore("https://s16.example/", [], TopicQuery("битва роботов"), since)
    urls = {p.url for p in report.pages}
    assert "https://s16.example/press/news/post/7" in urls            # со второй страницы списка
    assert "https://s16.example/press/news/post/3" not in urls        # 2024 год — раньше периода


def test_dead_site_is_not_retried_for_every_page(fake_web):
    report = SiteExplorer(HttpClient()).explore("https://down.example/", [], TopicQuery("битва"), None)
    assert report.unreachable and len([c for c in fake_web if "down.example" in c]) <= 2


@pytest.fixture
def schools(org):
    s16 = School.objects.create(organization=org, number=13, name="СОШ №16", website="https://s16.example/")
    l14 = School.objects.create(organization=org, number=12, name="Лицей №14")
    Source.objects.create(school=l14, kind="html", url="https://l14.example/news")
    down = School.objects.create(organization=org, number=1, name="СОШ №1", website="https://down.example/")
    nosite = School.objects.create(organization=org, number=2, name="СОШ №2")
    return s16, l14, down, nosite


@pytest.mark.django_db
def test_search_on_server_with_evidence_and_exports(org, admin_user, schools, fake_web, client):
    s16, l14, down, nosite = schools
    search = svc.start_search(org, admin_user, "битва роботов", since=date(2026, 1, 1))
    search.refresh_from_db()
    assert search.status == "done"
    res = {r.school_id: r for r in search.results.all()}
    assert res[s16.pk].status == "found" and res[s16.pk].hits_count == 1
    assert res[l14.pk].status == "none"
    assert res[down.pk].status == "unreachable"
    assert res[nosite.pk].status == "no_site"
    hit = res[s16.pk].hits.get()
    assert hit.published_at.date() == date(2026, 10, 12) and "битве роботов" in hit.snippet
    assert hit.screenshot and hit.html_copy and len(hit.checksum) == 64

    wb = load_workbook(io.BytesIO(svc.export_search_xlsx(search)))
    rows = {r[2]: r for r in wb["Школы"].iter_rows(min_row=3, values_only=True)}
    assert rows["СОШ №16"][4] == "Есть" and rows["Лицей №14"][4] == "Нет"
    assert rows["СОШ №1"][4] == "Сайт не открылся"
    with zipfile.ZipFile(svc.export_search_zip(search)) as z:
        names = z.namelist()
    assert "Итоги.xlsx" in names
    assert any(n.startswith("13 СОШ №16/2026-10-12 Наши победы") and n.endswith(".png") for n in names)

    client.force_login(admin_user)
    page = client.get(f"/app/search/{search.pk}/").content.decode()
    assert "Наши победы" in page and "Сайт не открылся" in page
    assert client.get(f"/app/search/{search.pk}/zip/").status_code == 200
    assert client.get("/app/search/").status_code == 200


@pytest.mark.django_db
def test_search_form_and_api(org, admin_user, schools, fake_web, client):
    client.force_login(admin_user)
    r = client.post("/app/search/", {"query": "битва роботов", "since": "2026-01-01", "screenshots": "on"})
    search = SchoolSearch.objects.get()
    assert r["Location"] == f"/app/search/{search.pk}/"
    data = client.get(f"/api/v1/searches/{search.pk}/").json()
    found = [x for x in data["results"] if x["status"] == "found"]
    assert found[0]["school"]["name"] == "СОШ №16" and found[0]["hits"][0]["screenshot"]
    r = client.post("/api/v1/searches/", {"query": " "}, content_type="application/json")
    assert r.status_code == 400


@pytest.mark.django_db
def test_search_through_agent(org, admin_user, schools, fake_web, client):
    from partners.agent import run_searches

    s16 = schools[0]
    token = CrawlAgent.issue(org)
    CrawlAgent.objects.update(last_seen_at=datetime.now(timezone.utc))
    search = svc.start_search(org, admin_user, "битва роботов", school_ids=[s16.pk])
    search.refresh_from_db()
    assert search.status == "agent"                     # сервер не обходит — ждём агента
    h = {"HTTP_AUTHORIZATION": f"Bearer {token}"}
    tasks = client.get("/api/v1/agent/tasks/", **h).json()["searches"]
    assert tasks[0]["site"] == "https://s16.example/"

    class Client:
        def search_result(self, data, files):
            payload = dict(data)
            payload.update({k: io.BytesIO(v[1]) for k, v in files.items()})
            for k, f in list(payload.items()):
                if isinstance(f, io.BytesIO):
                    f.name = files[k][0]
            r = client.post("/api/v1/agent/search-result/", payload, **h)
            assert r.status_code == 200, r.content
            return r.json()

    class Snap:
        def capture(self, url):
            return b"\xff\xd8\xff fake jpeg"

    out = run_searches(Client(), tasks, log=lambda *_: None, snapshotter=Snap())
    assert out == {"schools": 1, "found": 1}
    search.refresh_from_db()
    r = search.results.get()
    assert search.status == "done" and r.checked_by_agent and r.status == "found"
    hit = r.hits.get(url="https://s16.example/press/news/post/7")
    assert hit.screenshot.name.endswith(".jpg") and hit.html_copy


@pytest.mark.django_db
def test_report_zip_has_school_folders(client, org, admin_user):
    client.force_login(admin_user)
    r = client.get("/app/report/zip/")
    with zipfile.ZipFile(io.BytesIO(b"".join(r.streaming_content))) as z:
        assert "Отчёт по школам.xlsx" in z.namelist()



def test_path_pagination_keeps_section(monkeypatch):
    """«/press/news/page/2/» — публикации на этой странице всё равно лежат в «/press/news/post/…»."""
    site = {
        "https://p.example/": '<a href="/press/news/">Новости</a>',
        "https://p.example/press/news/": '<a href="/press/news/post/1/">Бал</a><a href="/press/news/page/2/">2</a>',
        "https://p.example/press/news/page/2/": '<a href="/press/news/post/2/">Экскурсия</a>',
        "https://p.example/press/news/post/1/": page("Бал", "Осенний бал."),
        "https://p.example/press/news/post/2/": page("Экскурсия", "Увидели битву роботов-сумоистов."),
    }

    def get(self, url):
        for u in (url, url + "/"):
            if u in site:
                return site[u]
        raise SourceUnavailable(f"{url}: сайт ответил кодом 404")

    monkeypatch.setattr(HttpClient, "get", get)
    q = TopicQuery("битва роботов")
    report = SiteExplorer(HttpClient()).explore("https://p.example/", [], q)
    assert [p.url for p in report.pages if q.match(p.title, p.text)] == ["https://p.example/press/news/post/2/"]
